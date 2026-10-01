"""Tests del descubrimiento de influencers — sin red (Graph y cookies mockeados)."""
from __future__ import annotations

from datetime import datetime, timedelta, timezone

import pytest

from src import business_discovery as bdm
from src import influencers as inf

AHORA = datetime(2026, 10, 1, 12, 0, tzinfo=timezone.utc)


@pytest.fixture()
def cx(tmp_path):
    conn = inf.connect(tmp_path / "influencers.db")
    yield conn
    conn.close()


def _ts(dias: int) -> str:
    return (AHORA - timedelta(days=dias)).strftime("%Y-%m-%dT%H:%M:%S+0000")


def _bd(handle, followers=20_000, dias=5, bio="", name="Mariana López", likes=900, n=12):
    return {"username": handle, "name": name, "biography": bio, "website": None,
            "followers_count": followers, "follows_count": 300, "media_count": 200,
            "media": {"data": [{"timestamp": _ts(dias + i), "like_count": likes,
                                "comments_count": 100} for i in range(n)]}}


# ---------------------------------------------------------------- DB

def test_connect_rechaza_gdlscene(tmp_path):
    with pytest.raises(RuntimeError):
        inf.connect(tmp_path / "gdlscene.db")


# ---------------------------------------------------------------- filtro barato

def test_filtro_descarta_privadas_y_vistas_sin_lookup():
    usuarios = [{"username": "pub", "is_private": False},
                {"username": "priv", "is_private": True},
                {"username": "Vista", "is_private": False},
                {"username": "pub", "is_private": False}]  # repetida en el listado
    pub, priv, ya = inf.filtrar_listado(usuarios, {"vista"})
    assert [u["username"] for u in pub] == ["pub"]
    assert [u["username"] for u in priv] == ["priv"]
    assert ya == 2


def test_guardar_listado_privadas_quedan_descartadas_y_no_se_reconsultan(cx):
    r = inf.guardar_listado(cx, "followers:cag_cucs", [
        {"username": "a", "is_private": True}, {"username": "b", "is_private": False}])
    assert r == {"listados": 2, "nuevas": 1, "privadas": 1, "ya_vistas": 0}
    fila = cx.execute("SELECT * FROM influencer_candidates WHERE handle='a'").fetchone()
    assert (fila["estado"], fila["motivo_descarte"], fila["fuente"]) == \
        ("descartada", "privada", "followers:cag_cucs")
    r2 = inf.guardar_listado(cx, "following:udg_oficial", [
        {"username": "a", "is_private": True}, {"username": "b", "is_private": False}])
    assert r2["ya_vistas"] == 2 and r2["nuevas"] == 0
    assert [p["handle"] for p in inf.pendientes(cx, 10)] == ["b"]


# ---------------------------------------------------------------- ER

def test_er_promedio_ultimos_12():
    medias = [{"like_count": 90, "comments_count": 10}] * 12 + \
             [{"like_count": 99999, "comments_count": 0}]  # el 13º no cuenta
    assert inf.calcular_er(medias, 10_000) == 1.0


def test_er_ignora_likes_ocultos_y_sin_followers():
    medias = [{"like_count": None, "comments_count": 5},
              {"like_count": 190, "comments_count": 10}]
    assert inf.calcular_er(medias, 1_000) == 20.0
    assert inf.calcular_er(medias, 0) is None
    assert inf.calcular_er([{"comments_count": 3}], 100) is None


# ---------------------------------------------------------------- criterio

@pytest.mark.parametrize("kw,esperado", [
    ({}, None),
    ({"is_private": True}, "privada"),
    ({"is_professional": False}, "no_profesional"),
    ({"followers": 9_999}, "<10k"),
    ({"followers": 2_000_000}, None),  # sin tope superior
    ({"ultimo": AHORA - timedelta(days=60)}, "inactiva"),
    ({"ultimo": None}, "inactiva"),
    ({"es_marca": True}, "marca"),
])
def test_criterio_candidata(kw, esperado):
    base = {"is_professional": True, "is_private": False, "followers": 10_000,
            "ultimo": AHORA - timedelta(days=59), "es_marca": False}
    base.update(kw)
    assert inf.motivo_no_candidata(**base, ahora=AHORA) == esperado


def test_tier():
    assert inf.tier_de(99_999) == "micro"
    assert inf.tier_de(100_000) == "macro"
    assert inf.tier_de(500) == "nano"


# ---------------------------------------------------------------- heurística

def test_marca_por_categoria_y_persona_gana():
    assert inf.es_marca("Restaurant", "", "Tacos Chuy")
    assert inf.es_marca("College & university", "", "CUCS")
    assert inf.es_marca("Media/news company", "", "")
    assert not inf.es_marca("Digital creator", "envíos a todo méxico, pedidos por DM", "")
    assert not inf.es_marca("Personal blog", "", "Fer")


def test_marca_por_texto_sin_categoria():
    assert inf.es_marca(None, "Envíos a todo México · Pedidos por DM · Sucursal Chapultepec", "")
    assert not inf.es_marca(None, "estudiante de medicina 🩺 gdl", "Fernanda")


def test_genero():
    assert inf.estimar_genero("Mariana López", "mamá de dos")[0] == "F"
    assert inf.estimar_genero("Mariana López", "mamá de dos")[1] in ("media", "alta")
    assert inf.estimar_genero("Carlos Ruiz", "papá, ingeniero")[0] == "M"
    assert inf.estimar_genero("", "", handle="xyz_123") == ("?", "baja")
    assert inf.estimar_genero("", "creadora de contenido y estudiante de psicología")[0] == "F"


def test_edad():
    assert inf.estimar_edad("Estudiante de psicología @cucs")[0] == "18-24"
    assert inf.estimar_edad("mamá de 2 · esposa · emprendedora")[0] == "25-40"
    assert inf.estimar_edad("22 años, gdl") == ("18-24", "alta")
    assert inf.estimar_edad("") == ("?", "baja")


def test_mx():
    assert inf.estimar_mx("Fer", "GDL, Jalisco 🇲🇽") == ("si", "alta")
    assert inf.estimar_mx("Sofi", "Buenos Aires 🇦🇷")[0] == "no"
    assert inf.estimar_mx("Ana", "amo la vida y el café, contacto para colaboraciones")[0] == "si"
    assert inf.estimar_mx("Ana", "", website="https://ana.com.mx") [0] == "si"
    assert inf.estimar_mx("", "") == ("?", "baja")


# ---------------------------------------------------------------- aplicar / lookup

def test_aplicar_perfil_candidata_y_descartes(cx):
    inf.guardar_listado(cx, "following:udg_oficial", [
        {"username": h, "is_private": False} for h in ("ok", "chica", "vieja", "taqueria")])
    assert inf.aplicar_perfil(cx, "ok", _bd("ok", followers=150_000, bio="GDL 🇲🇽"),
                              "Digital creator", ahora=AHORA) == "candidata"
    fila = cx.execute("SELECT * FROM influencer_candidates WHERE handle='ok'").fetchone()
    assert fila["tier"] == "macro" and fila["er_pct"] == round(1000 / 150_000 * 100, 2)
    assert fila["mx_est"] == "si" and fila["genero_est"] == "F"
    inf.aplicar_perfil(cx, "chica", _bd("chica", followers=5_000), None, ahora=AHORA)
    inf.aplicar_perfil(cx, "vieja", _bd("vieja", dias=90), None, ahora=AHORA)
    inf.aplicar_perfil(cx, "taqueria", _bd("taqueria"), "Mexican restaurant", ahora=AHORA)
    motivos = dict(cx.execute("SELECT handle, motivo_descarte FROM influencer_candidates"
                              " WHERE estado='descartada'").fetchall())
    assert motivos == {"chica": "<10k", "vieja": "inactiva", "taqueria": "marca"}


def test_lookup_respeta_budget_y_para_en_rate_limit(cx, monkeypatch):
    monkeypatch.setattr(inf, "ultimo_post", lambda m: datetime.now(timezone.utc))
    inf.guardar_listado(cx, "followers:x", [
        {"username": f"u{i}", "is_private": False} for i in range(6)])
    llamadas = []

    def fetch(handle, **_):
        llamadas.append(handle)
        if handle == "u0":
            raise bdm.NoEsBusiness("(#110) requires Business or Creator")
        if handle == "u1":
            return _bd(handle)
        raise bdm.GraphRateLimited("(#4) rate limit")

    res = inf.lookup(cx, 4, _fetch=fetch, _categoria=lambda h: "Digital creator", _pausa=0)
    assert llamadas == ["u0", "u1", "u2"]  # paró en u2, no siguió
    assert res["no_profesional"] == 1 and res["candidatas"] == 1
    assert res["parada"].startswith("GraphRateLimited")
    assert len(inf.pendientes(cx, 99)) == 4  # u2..u5 siguen pendientes → reanudable


def test_lookup_categoria_solo_para_las_que_pasan_umbral(cx, monkeypatch):
    monkeypatch.setattr(inf, "ultimo_post", lambda m: datetime.now(timezone.utc))
    inf.guardar_listado(cx, "followers:x", [
        {"username": h, "is_private": False} for h in ("grande", "chica")])
    pedidas = []

    def cat(h):
        pedidas.append(h)
        return "Digital creator"

    fetch = lambda h, **_: _bd(h, followers=50_000 if h == "grande" else 800)  # noqa: E731
    inf.lookup(cx, 10, _fetch=fetch, _categoria=cat, _pausa=0)
    assert pedidas == ["grande"]


def test_lookup_para_si_se_quema_la_cookie(cx, monkeypatch):
    monkeypatch.setattr(inf, "ultimo_post", lambda m: datetime.now(timezone.utc))
    inf.guardar_listado(cx, "followers:x", [
        {"username": h, "is_private": False} for h in ("a", "b")])

    def cat(h):
        raise inf.SesionQuemada("cuenta scraper 'x' quemada: HTTP 429")

    res = inf.lookup(cx, 10, _fetch=lambda h, **_: _bd(h), _categoria=cat, _pausa=0)
    assert res["llamadas"] == 1 and "429" in res["parada"]


# ---------------------------------------------------------------- export

def _sembrar_candidatas(cx):
    inf.guardar_listado(cx, "following:udg_oficial", [
        {"username": h, "is_private": False} for h in ("baja", "alta")])
    inf.aplicar_perfil(cx, "baja", _bd("baja", likes=100), "Blogger", ahora=AHORA)
    inf.aplicar_perfil(cx, "alta", _bd("alta", likes=5000), "Blogger", ahora=AHORA)


def test_export_csv_ordenado_por_er(cx, tmp_path):
    _sembrar_candidatas(cx)
    out = tmp_path / "c.csv"
    assert inf.export_csv(cx, out) == 2
    lineas = out.read_text().splitlines()
    assert lineas[1].startswith("alta,https://www.instagram.com/alta/")
    assert lineas[2].startswith("baja,")


def test_export_xlsx(cx, tmp_path):
    from openpyxl import load_workbook
    _sembrar_candidatas(cx)
    out = tmp_path / "c.xlsx"
    inf.export_xlsx(cx, out)
    wb = load_workbook(out)
    assert wb.sheetnames == ["Candidatas", "Resumen"]
    ws = wb["Candidatas"]
    enc = [c.value for c in ws[1]]
    assert enc[-2:] == ["estado_contacto", "notas"]
    assert ws.freeze_panes == "A2" and ws.auto_filter.ref.startswith("A1:")
    assert ws["A2"].value == "alta"
    assert ws["A2"].hyperlink.target == "https://www.instagram.com/alta/"
    assert ws.cell(row=2, column=enc.index("followers") + 1).number_format == "#,##0"
    assert ws.cell(row=2, column=enc.index("er_pct") + 1).number_format == "0.0"
    rs = wb["Resumen"]
    assert rs["A1"].value == "fuente" and rs["A2"].value == "following:udg_oficial"


def test_embudo_por_fuente(cx):
    inf.guardar_listado(cx, "followers:cag_cucs", [
        {"username": "p", "is_private": True}, {"username": "q", "is_private": False}])
    inf._descartar(cx, "q", "no_profesional", is_professional=0)
    e = inf.embudo(cx)[0]
    assert (e["listados"], e["privadas"], e["consultadas"], e["no_profesionales"]) == \
        (2, 1, 1, 1)


# ---------------------------------------------------------------- followers

def test_listar_followers_respeta_tope(monkeypatch):
    from src import import_followees as imf
    monkeypatch.setattr(imf.time, "sleep", lambda s: None)
    pedidas = []

    class R:
        status_code = 200

        def __init__(self, i):
            self.i = i

        def raise_for_status(self):
            pass

        def json(self):
            return {"users": [{"username": f"u{self.i}_{k}"} for k in range(24)],
                    "next_max_id": str(self.i + 1)}

    class S:
        def get(self, url, params, timeout):
            pedidas.append((url, params.get("max_id")))
            return R(len(pedidas))

    users = imf.listar_followers(S(), "123", limite=50)
    assert len(users) == 50 and len(pedidas) == 3
    assert "/friendships/123/followers/" in pedidas[0][0]


def test_listar_followers_401_es_rate_limit():
    from src import import_followees as imf
    from src.ingest_ig import IngestRateLimited

    class S:
        def get(self, *a, **k):
            return type("R", (), {"status_code": 401})()

    with pytest.raises(IngestRateLimited):
        imf.listar_followers(S(), "1", 10)
